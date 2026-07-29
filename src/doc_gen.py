"""文档生成工具 - 使用python-docx生成Word文档（替代coze_coding_dev_sdk）"""
import os
import re
import datetime
import logging
from typing import Optional, List
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

logger = logging.getLogger(__name__)

DATA_DIR = os.environ.get("DATA_DIR", "/data")
REPORTS_DIR = os.path.join(DATA_DIR, "reports")
os.makedirs(REPORTS_DIR, exist_ok=True)


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    """添加标题"""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = "Microsoft YaHei"
        run.font.color.rgb = RGBColor(0x1e, 0x40, 0xaf) if level == 1 else RGBColor(0x33, 0x33, 0x33)


def _add_paragraph(doc: Document, text: str, bold: bool = False, italic: bool = False,
                   color: Optional[tuple] = None, font_size: int = 11) -> None:
    """添加段落"""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def _add_separator(doc: Document) -> None:
    """添加分隔线"""
    p = doc.add_paragraph()
    run = p.add_run("─" * 60)
    run.font.color.rgb = RGBColor(0xcc, 0xcc, 0xcc)
    run.font.size = Pt(8)


def _parse_inline_formatting(p, text: str) -> None:
    """解析简单行内格式：**bold** *italic* [text](url)"""
    # 简化处理：去掉markdown标记，显示纯文本
    clean = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    clean = re.sub(r'\*(.*?)\*', r'\1', clean)
    clean = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', clean)
    run = p.add_run(clean)
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(11)


def generate_docx_from_markdown(markdown_content: str, title: str) -> str:
    """从Markdown内容生成Word文档，返回本地文件路径"""
    doc = Document()

    # 设置默认字体
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(11)

    lines: List[str] = markdown_content.split("\n")
    i: int = 0
    in_table: bool = False
    table_rows: List[List[str]] = []

    while i < len(lines):
        line: str = lines[i].strip()

        # 空行
        if not line:
            i += 1
            continue

        # 分隔线
        if line.startswith("---"):
            _add_separator(doc)
            i += 1
            continue

        # H1标题
        if line.startswith("# ") and not line.startswith("## "):
            text = line[2:].strip()
            _add_heading(doc, text, level=1)
            i += 1
            continue

        # H2标题
        if line.startswith("## ") and not line.startswith("### "):
            text = line[3:].strip()
            _add_heading(doc, text, level=2)
            i += 1
            continue

        # H3标题
        if line.startswith("### "):
            text = line[4:].strip()
            _add_heading(doc, text, level=3)
            i += 1
            continue

        # 引用块
        if line.startswith(">"):
            text = line.lstrip("> ").strip()
            # 收集多行引用
            quote_lines: List[str] = [text]
            j = i + 1
            while j < len(lines) and lines[j].strip().startswith(">"):
                quote_lines.append(lines[j].strip().lstrip("> "))
                j += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            run = p.add_run("\n".join(quote_lines))
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            run.italic = True
            i = j
            continue

        # 表格（简单处理）
        if "|" in line and ("---" in lines[i + 1] if i + 1 < len(lines) else False):
            table_rows = []
            header = [c.strip() for c in line.strip("|").split("|")]
            table_rows.append(header)
            i += 2  # 跳过分隔行
            while i < len(lines) and "|" in lines[i]:
                row = [c.strip() for c in lines[i].strip("|").split("|")]
                table_rows.append(row)
                i += 1
            if table_rows:
                tbl = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
                tbl.style = "Light Grid Accent 1"
                for ri, row in enumerate(table_rows):
                    for ci, cell in enumerate(row):
                        if ci < len(tbl.rows[ri].cells):
                            tbl.rows[ri].cells[ci].text = cell.replace("**", "").replace("*", "")
                            for p in tbl.rows[ri].cells[ci].paragraphs:
                                for run in p.runs:
                                    run.font.name = "Microsoft YaHei"
                                    run.font.size = Pt(10)
                                    if ri == 0:
                                        run.bold = True
            continue

        # 有序列表
        list_match = re.match(r'^(\d+)\.\s+(.*)', line)
        if list_match:
            num, text = list_match.groups()
            p = doc.add_paragraph(style="List Number")
            _parse_inline_formatting(p, text)
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        _parse_inline_formatting(p, line)
        i += 1

    # 保存文件
    now = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    filename: str = f"{title}_{now}.docx"
    filepath: str = os.path.join(REPORTS_DIR, filename)
    doc.save(filepath)
    logger.info(f"Word文档已生成: {filepath}")
    return filepath


def generate_report_html(title: str, content: str) -> str:
    """生成HTML报告（备用），返回文件路径"""
    now = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    filename: str = f"{title}_{now}.html"
    filepath: str = os.path.join(REPORTS_DIR, filename)
    
    html = f"""<!DOCTYPE html>
<html>
<head>
    <meta charset="UTF-8">
    <title>{title}</title>
    <style>
        body {{ font-family: "Microsoft YaHei", "SimHei", sans-serif; max-width: 900px; margin: 40px auto; padding: 20px; line-height: 1.8; color: #333; }}
        h1 {{ color: #1e40af; border-bottom: 3px solid #1e40af; padding-bottom: 10px; }}
        h2 {{ color: #333; border-left: 4px solid #1e40af; padding-left: 12px; margin-top: 30px; }}
        h3 {{ color: #555; }}
        blockquote {{ border-left: 4px solid #ddd; margin: 15px 0; padding: 10px 20px; background: #f9f9f9; color: #666; }}
        a {{ color: #1e40af; text-decoration: none; }}
        a:hover {{ text-decoration: underline; }}
        hr {{ border: none; border-top: 1px solid #ddd; margin: 30px 0; }}
        .meta {{ color: #888; font-size: 12px; }}
    </style>
</head>
<body>
{content}
</body>
</html>"""
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(html)
    return filepath
